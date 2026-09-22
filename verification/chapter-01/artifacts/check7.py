# 원고 03-chapter-01.md :259-265 확인 목록 7개를 구조에 의존하지 않고 검사한다.
import sys, re, os
import openpyxl

path = sys.argv[1]
res = []
def ok(n, label, passed, note=""):
    res.append((n, label, passed, note))

# 1. 파일이 실제로 생겼다
ok(1, "movie-report-2024.xlsx 파일이 실제로 생겼다", os.path.basename(path) == "movie-report-2024.xlsx" and os.path.exists(path),
   "%d bytes" % os.path.getsize(path) if os.path.exists(path) else "없음")

# 2. 정상적으로 열린다
try:
    wb = openpyxl.load_workbook(path)
    ok(2, "파일이 정상적으로 열린다", True, "시트 %s" % wb.sheetnames)
except Exception as e:
    ok(2, "파일이 정상적으로 열린다", False, str(e)); wb = None

cells = []   # (sheet, row, col, value)
if wb:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    cells.append((ws.title, c.row, c.column, c.value))
texts = [str(v) for _,_,_,v in cells]
blob = "\n".join(texts)

# 3. 영화 TOP 20이 보인다 — 순위 1..20이 한 열에 연속하고, 그 옆에 영화명 20개
top20 = False; t20note = ""
if wb:
    for ws in wb.worksheets:
        for col in range(1, ws.max_column+1):
            seq = [ws.cell(r, col).value for r in range(1, ws.max_row+1)]
            for start in range(len(seq)-19):
                window = seq[start:start+20]
                if window == list(range(1,21)):
                    names = [ws.cell(start+1+i+0, col+1).value for i in range(20)]
                    names = [n for n in names if isinstance(n, str) and n.strip()]
                    if len(names) == 20:
                        top20 = True
                        t20note = "시트 '%s' %d~%d행, 순위=%d열 영화명=%d열" % (
                            ws.title, start+1, start+20, col, col+1)
                    break
            if top20: break
        if top20: break
ok(3, "영화 TOP 20이 보인다", top20, t20note)

# 4. 관객수와 매출액이 보인다 — 머리글 문구 + 20행 정수 데이터
def numeric_col_under(label_re):
    if not wb: return None
    for ws in wb.worksheets:
        for s,r,c,v in [x for x in cells if x[0]==ws.title]:
            if isinstance(v,str) and re.search(label_re, v):
                nums = [ws.cell(r+i, c).value for i in range(1,21)]
                nums = [n for n in nums if isinstance(n,(int,float))]
                if len(nums) == 20:
                    return "시트 '%s' 머리글 %r (%d행 %d열), 데이터 %d~%d행" % (ws.title, v, r, c, r+1, r+20)
    return None
aud = numeric_col_under(r"관객수")
sal = numeric_col_under(r"매출액")
ok(4, "관객수와 매출액이 보인다", bool(aud and sal), "관객수: %s / 매출액: %s" % (aud, sal))

# 5. 핵심 요약이 있다 — 합계·평균 같은 집계값이 표 밖에 있고 실제 합계와 일치
summary = False; snote = ""
if aud:
    m = re.search(r"시트 '(.+?)' 머리글 .+? \((\d+)행 (\d+)열\)", aud)
    ws = wb[m.group(1)]; r0=int(m.group(2)); c0=int(m.group(3))
    total = sum(ws.cell(r0+i, c0).value for i in range(1,21))
    hit = []
    for s,r,c,v in cells:
        if isinstance(v,(int,float)) and abs(v-total) < 1:
            if not (s==ws.title and c==c0 and r0 < r <= r0+20): hit.append((s,r,c,"값"))
        if isinstance(v,str):
            digits = re.sub(r"[^\d]", "", v)
            if digits and digits == str(total): hit.append((s,r,c,"문자열"))
    summary = bool(hit)
    snote = "관객수 합계 %d 이 요약에 등장: %s" % (total, hit[:3])
ok(5, "핵심 요약이 있다(요약값이 상세 20행 합계와 일치)", summary, snote)

# 6. 실제 데이터를 이용한 차트가 있다 — 차트 객체가 있고 셀 참조를 가진다
charts = []; refs = []
if wb:
    for ws in wb.worksheets:
        for ch in getattr(ws, "_charts", []):
            charts.append((ws.title, type(ch).__name__))
            for ser in getattr(ch, "series", []):
                try: refs.append(ser.val.numRef.f)
                except Exception: pass
ok(6, "실제 데이터를 이용한 차트가 있다", bool(charts) and bool(refs), "차트 %s / 참조 %s" % (charts, refs))

# 7. KOBIS 출처·조회 기준 연도 2024·재현 정보
has_src = bool(re.search(r"KOBIS|영화진흥위원회", blob))
has_year = bool(re.search(r"2024", blob))
has_repro = bool(re.search(r"kobis\.or\.kr", blob))
ok(7, "KOBIS 출처·조회 기준 연도 2024·다시 찾을 수 있는 정보", has_src and has_year and has_repro,
   "출처=%s 연도=%s 주소=%s" % (has_src, has_year, has_repro))

npass = sum(1 for _,_,p,_ in res if p)
print("## %s" % path)
for n,label,p,note in res:
    print("%d. [%s] %s" % (n, "PASS" if p else "FAIL", label))
    if note: print("     %s" % note)
print("=> %d/7 PASS" % npass)
sys.exit(0 if npass==7 else 1)
