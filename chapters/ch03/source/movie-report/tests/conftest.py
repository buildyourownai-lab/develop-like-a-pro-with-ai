"""movie-report-2024.xlsx 검사에서 공통으로 쓰는 파일 로딩 도우미들."""

from pathlib import Path

import openpyxl
import pytest

ROOT_DIR = Path(__file__).resolve().parent.parent
REPORT_PATH = ROOT_DIR / "movie-report-2024.xlsx"
SOURCE_PATH = ROOT_DIR / "source" / "kobis-2024-source.xlsx"

TOP20_HEADER_ROW = 3
TOP20_ROW_COUNT = 20


def _read_top20_rows(sheet):
    """TOP20 표 시트에서 (순위, 영화명, 개봉일, 관객수, 매출액) 20행을 읽어온다."""
    rows = []
    for i in range(TOP20_ROW_COUNT):
        r = TOP20_HEADER_ROW + 1 + i
        rows.append(
            {
                "순위": sheet.cell(row=r, column=1).value,
                "영화명": sheet.cell(row=r, column=2).value,
                "개봉일": sheet.cell(row=r, column=3).value,
                "관객수": sheet.cell(row=r, column=4).value,
                "매출액": sheet.cell(row=r, column=5).value,
            }
        )
    return rows


def _read_info_sheet(sheet):
    """'항목' / '값'(또는 '내용') 2열 표를 {항목: 값} 딕셔너리로 읽어온다."""
    info = {}
    for row in sheet.iter_rows(values_only=True):
        if row[0] in (None, "항목") or row[1] is None:
            continue
        info[row[0]] = row[1]
    return info


@pytest.fixture(scope="session")
def report_wb():
    assert REPORT_PATH.exists(), f"검사 대상 파일이 없습니다: {REPORT_PATH}"
    return openpyxl.load_workbook(REPORT_PATH, data_only=True)


@pytest.fixture(scope="session")
def source_wb():
    assert SOURCE_PATH.exists(), f"보관해 둔 기준 자료 파일이 없습니다: {SOURCE_PATH}"
    return openpyxl.load_workbook(SOURCE_PATH, data_only=True)


@pytest.fixture(scope="session")
def report_top20(report_wb):
    return _read_top20_rows(report_wb["TOP20"])


@pytest.fixture(scope="session")
def source_top20(source_wb):
    return _read_top20_rows(source_wb["원본데이터"])


@pytest.fixture(scope="session")
def report_summary(report_wb):
    return _read_info_sheet(report_wb["요약"])


@pytest.fixture(scope="session")
def report_source_info(report_wb):
    return _read_info_sheet(report_wb["출처"])
